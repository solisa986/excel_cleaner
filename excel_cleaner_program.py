import openpyxl
from pathlib import Path
import pandas as pd
import streamlit as st

workbook_name = input("Please enter the name of the workbook (Note: only accepts .xlsx files): ").replace(".xlsx", "")
# Load the Excel file
workbook = openpyxl.load_workbook(f'C:/Users/agarcia/Scripts - Adriana/Uncleaned Excel Files/{workbook_name}.xlsx')

sheet_name = input("Please enter the name of the sheet that needs to be cleaned: ")

# Select the worksheet by name
worksheet = workbook[f'{sheet_name}']

table_list = list(worksheet.tables.keys())
print(table_list)

extracted_table = []
for t in table_list:
    lookup_table = worksheet.tables[f'{t}']
    lookup_range = lookup_table.ref

    # Access the data in the table range
    data = worksheet[lookup_range]
    rows_list = []

    # Loop through each row and get the values in the cells
    for row in data:
        # Get a list of all columns in each row
        cols = []
        for col in row:
            cols.append(col.value)
        rows_list.append(cols)

    # Create a pandas dataframe from the rows_list.
    # The first row is the column names
    df = pd.DataFrame(data=rows_list[1:], index=None, columns=rows_list[0])
    print("Found the following columns in your table.")
    print("Please ")
    filtered_df = df.filter(items=[])
    print(df)